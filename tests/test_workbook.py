"""The Excel model, recalculated with the `formulas` engine, agrees with the Python reference model.

The workbook is built from the fixture store, compiled once, then recalculated under each
scenario by overriding the scenario selector cell.
"""
from __future__ import annotations

import re

import openpyxl
import pandas as pd
import pytest

from saas_metrics import drivers, model, workbook

formulas = pytest.importorskip("formulas")
NAME = "saas_operating_model.xlsx"
MODEL_SHEETS = ("Cohort_ARR_Build", "Headcount_Opex_Build", "Operating_Model")


@pytest.fixture(scope="module")
def built(db, tmp_path_factory):
    path = tmp_path_factory.mktemp("wb") / NAME
    workbook.build(db, path)
    return path


@pytest.fixture(scope="module")
def layout(built):
    """sheet -> {row key: row number}, read from column A, and the quarter columns."""
    wb = openpyxl.load_workbook(built)
    rows = {s: {wb[s].cell(r, 1).value: r for r in range(workbook.FIRST_ROW, wb[s].max_row + 1)
                if wb[s].cell(r, 1).value} for s in MODEL_SHEETS}
    ws = wb["Operating_Model"]
    cols = {ws.cell(workbook.HDR["period_end"], c).value: c
            for c in range(workbook.FIRST_COL, ws.max_column + 1) if ws.cell(workbook.HDR["period_end"], c).value}
    return wb, rows, cols


@pytest.fixture(scope="module")
def compiled(built):
    return formulas.ExcelModel().loads(str(built)).finish()


def _solve(compiled, scenario):
    sol = compiled.calculate(inputs={f"'[{NAME}]INPUTS_SCENARIOS'!C3": scenario})
    out = {}
    for k, v in sol.items():
        m = re.match(r"'\[[^\]]+\]([^']+)'!([A-Z]+\d+)$", k)
        if m:
            val = v.value[0][0] if hasattr(v, "value") else v
            out[(m.group(1), m.group(2))] = val
    return out


@pytest.fixture(scope="module")
def reference(db):
    hist = drivers.history(db)
    cal = db.execute("SELECT * FROM cohort_calibration WHERE vintage='as_of_cutoff'").df().set_index("period_end")
    hc = drivers.headcount_history(db, hist).set_index("period_end")["headcount"]
    drv = drivers.derive(db)
    return {sc: model.run(hist, cal, hc, drv, sc) for sc in drivers.SCENARIOS}


# workbook (sheet, row key) -> reference model column
COMPARE = {
    **{("Operating_Model", k): k for k in (
        "atlas", "ea", "svc", "sub", "revenue", "cost_sub", "cost_svc", "cost", "gp", "rnd", "snm", "gna",
        "opex", "op_inc", "sbc_total", "amort", "ng_op_inc", "fcf",
        *(f"sbc_{s}" for s in model.SBC_LINES))},
    ("Cohort_ARR_Build", "customers"): "customers",
    ("Cohort_ARR_Build", "arr"): "arr",
    ("Cohort_ARR_Build", "ea_ttm"): "ea_ttm",
    ("Headcount_Opex_Build", "headcount"): "headcount",
    **{("Headcount_Opex_Build", f"ng_{d}"): f"ng_{d}" for d in drivers.DEPTS},
    **{("Headcount_Opex_Build", f"cph_{d}"): f"cph_{d}" for d in drivers.DEPTS},
}


@pytest.mark.parametrize("scenario", drivers.SCENARIOS)
def test_workbook_matches_reference_model(compiled, layout, reference, scenario):
    _, rows, cols = layout
    sol = _solve(compiled, scenario)
    ref = reference[scenario]
    compared, bad = 0, []
    for (sheet, key), col in COMPARE.items():
        for period_end, c in cols.items():
            expected = ref.at[period_end, col] if period_end in ref.index else None
            if expected is None or pd.isna(expected):
                continue
            addr = f"{openpyxl.utils.get_column_letter(c)}{rows[sheet][key]}"
            got = sol.get((sheet.upper(), addr))
            compared += 1
            if not isinstance(got, (int, float)) or abs(got - expected) > 1e-6 * max(1.0, abs(expected)):
                bad.append((sheet, key, period_end, got, expected))
    assert compared > 600, compared
    assert not bad, bad[:5]


@pytest.mark.parametrize("scenario", drivers.SCENARIOS)
def test_in_workbook_checks_pass(compiled, scenario):
    sol = _solve(compiled, scenario)
    results = {a: v for (s, a), v in sol.items() if s == "CHECKS" and a.startswith("C") and a != "C3"}
    assert len(results) == 6 and set(results.values()) == {"PASS"}, results


def test_scenarios_differ_and_are_ordered(reference):
    rev = {sc: reference[sc][reference[sc].kind == "Forecast"].revenue.sum() for sc in drivers.SCENARIOS}
    assert rev["Downside"] < rev["Base"] < rev["Upside"]


def test_no_numeric_constants_in_formulas(layout):
    """Only 0 and 1 may appear as literals (MATCH type, 1 + growth); 365 and 4 are input cells."""
    wb, _, _ = layout
    offenders = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    body = re.sub(r"('[^']+'|[A-Za-z_]+)?!?\$?[A-Z]{1,3}\$?\d+", "", c.value)  # drop references
                    body = re.sub(r'"[^"]*"', "", body)                                       # drop strings
                    lits = [x for x in re.findall(r"\d+(?:\.\d+)?", body) if x not in ("0", "1")]
                    if lits:
                        offenders.append((ws.title, c.coordinate, c.value))
    assert not offenders, offenders[:5]


def test_forecast_cells_are_formulas(layout):
    wb, rows, cols = layout
    fc_cols = [c for p, c in cols.items() if p > "2025-07-31"]
    assert len(fc_cols) == 4
    for sheet in MODEL_SHEETS:
        for key, r in rows[sheet].items():
            for c in fc_cols:
                v = wb[sheet].cell(r, c).value
                assert v is None or (isinstance(v, str) and v.startswith("=")), (sheet, key, c, v)


def test_history_inputs_tie_to_sec_as_of_cutoff(db, layout):
    """Every blue historical input equals the as-of-cutoff SEC value for that line and quarter."""
    wb, rows, cols = layout
    q = db.execute("SELECT period_end, line, value FROM quarterly_financials WHERE vintage='as_of_cutoff'").df()
    sec = {(r.line, r.period_end): r.value for r in q.itertuples()}
    om = wb["Operating_Model"]
    pairs = {"atlas": "revenue_atlas", "ea": "revenue_other_subscription", "svc": "revenue_services",
             "cost_sub": "cost_subscription", "cost_svc": "cost_services", "rnd": "rnd", "snm": "snm",
             "gna": "gna", "op_inc_reported": "operating_income", "cfo": "cfo", "capex": "capex",
             "amort": "amort_intangibles", **{f"sbc_{s}": f"sbc_{s}" for s in model.SBC_LINES}}
    n = 0
    for key, line in pairs.items():
        for p, c in cols.items():
            if p > "2025-07-31":
                continue
            assert om.cell(rows["Operating_Model"][key], c).value == sec[(line, p)], (key, p)
            n += 1
    assert n == 15 * len(pairs)
