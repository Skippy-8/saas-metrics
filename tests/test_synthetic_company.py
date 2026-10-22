"""The full model runs on a second, differently shaped company: SynthCo (FICTIONAL, configs/examples).

Four revenue segments (two on a trailing-12-month ARR basis), a December fiscal year end and its
own KPI names, built from generated filing-style tables (scripts/make_synthetic_company.py).
Proves the operating model, cohort ledger, workbook and backtest follow the config rather than
MongoDB's layout: QC passes, the recalculated workbook equals the Python model, the bridge
reconciles and the workbook bridge equals the Python bridge.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd
import pytest

from saas_metrics import backtest, build, config, drivers, model, qc, segments, workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import make_synthetic_company as synth  # noqa: E402

formulas = pytest.importorskip("formulas")
NAME = "synthco_model.xlsx"


@pytest.fixture(scope="module")
def synth_db():
    """The built store; SynthCo stays the active company until this module's tests finish."""
    with config.using(synth.CONFIG):
        con = duckdb.connect()
        build.assemble(con, synth.source_tables())
        yield con


@pytest.fixture(scope="module")
def built(synth_db, tmp_path_factory):
    path = tmp_path_factory.mktemp("synth") / NAME
    workbook.build(synth_db, path)
    wb = openpyxl.load_workbook(path)
    rows = {ws.title: {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
            for ws in wb.worksheets}
    om = wb["Operating_Model"]
    cols = {om.cell(workbook.HDR["period_end"], c).value: c for c in range(workbook.FIRST_COL, om.max_column + 1)
            if om.cell(workbook.HDR["period_end"], c).value}
    return path, wb, rows, cols


@pytest.fixture(scope="module")
def compiled(built):
    return formulas.ExcelModel().loads(str(built[0])).finish()


def _solve(compiled, scenario: str) -> dict:
    out = {}
    for k, v in compiled.calculate(inputs={f"'[{NAME}]INPUTS_SCENARIOS'!C3": scenario}).items():
        m = re.match(r"'\[[^\]]+\]([^']+)'!([A-Z]+\d+)$", k)
        if m:
            out[(m.group(1), m.group(2))] = v.value[0][0] if hasattr(v, "value") else v
    return out


def test_config_is_the_active_shape(synth_db):
    assert config.TICKER == "SYNTHCO" and config.ACTIVE.fiscal_year_end_month == 12
    assert [s.key for s in segments.active()] == ["platform", "addons", "licenses", "prosvc"]
    assert [s.key for s in segments.ttm()] == ["addons", "licenses"]
    assert list(backtest.GROUPS[3:6]) == ["addons_mix_timing", "licenses_mix_timing", "prosvc_mix_timing"]
    assert backtest.EXPLAINS == {"revenue": 6, "ng_op_inc": 10, "fcf": 11}
    assert [model.periods.label(e) for e in model.forecast_quarters()] == ["Q3 FY25", "Q4 FY25", "Q1 FY26", "Q2 FY26"]


def test_qc_passes(synth_db):
    checks = [qc.completeness(synth_db), *qc.identities(synth_db), qc.quarters_sum_to_year(synth_db),
              qc.fact_conflicts(synth_db), qc.point_in_time(synth_db), *qc.cohort_ledger(synth_db)]
    failed = [f"{c.name}: {c.detail}" for c in checks if not c.passed]
    assert len(checks) == 22 and not failed, failed
    # December year end: Q1..Q4 sum to the 10-K year (this check once assumed a February 1 fiscal year start)
    years = next(c for c in checks if c.name == "derived quarters sum to the 10-K fiscal year")
    assert years.detail.startswith("72 line-years"), years.detail   # 24 duration lines x FY23-FY25


def test_quarters_use_every_derivation_path(synth_db):
    d = synth_db.execute("SELECT DISTINCT derivation FROM quarterly_financials").df().derivation
    assert set(d) == {"reported_3m", "ytd_diff", "fy_minus_9m"}


@pytest.mark.parametrize("scenario", drivers.SCENARIOS)
def test_workbook_matches_reference_model(synth_db, built, compiled, scenario):
    _, _, rows, cols = built
    hist = drivers.history(synth_db)
    cal = synth_db.execute("SELECT * FROM cohort_calibration WHERE vintage='as_of_cutoff'").df().set_index("period_end")
    hc = drivers.headcount_history(synth_db, hist).set_index("period_end")["headcount"]
    ref = model.run(hist, cal, hc, drivers.derive(synth_db), scenario)
    keys = [s.key for s in segments.active()] + ["sub", "revenue", "cost", "gp", "opex", "op_inc", "sbc_total",
                                                   "ng_op_inc", "fcf"]
    compare = {("Operating_Model", k): k for k in keys}
    compare.update({("Cohort_ARR_Build", "arr"): "arr", ("Cohort_ARR_Build", "customers"): "customers",
                    **{("Cohort_ARR_Build", f"{s.key}_ttm"): f"{s.key}_ttm" for s in segments.ttm()}})
    sol = _solve(compiled, scenario)
    n, bad = 0, []
    for (sheet, key), col in compare.items():
        for period_end, c in cols.items():
            exp = ref.at[period_end, col] if period_end in ref.index else None
            if exp is None or pd.isna(exp):
                continue
            got = sol.get((sheet.upper(), f"{openpyxl.utils.get_column_letter(c)}{rows[sheet][key]}"))
            n += 1
            if not isinstance(got, (int, float)) or abs(got - exp) > 1e-6 * max(1.0, abs(exp)):
                bad.append((sheet, key, period_end, got, exp))
    assert n > 250 and not bad, bad[:5]
    checks = {a: v for (s, a), v in sol.items() if s == "CHECKS" and a.startswith("C") and a != "C3"}
    assert len(checks) == 6 and set(checks.values()) == {"PASS"}, checks


def test_bridge_reconciles_and_workbook_agrees(synth_db, built, compiled):
    _, _, rows, cols = built
    long, _ = backtest.bridge(synth_db)
    for metric, n in backtest.EXPLAINS.items():
        t = long[long.metric == metric].pivot(index="component", columns="period_end", values="value")
        resid = t.loc["plan"] + t.loc[list(backtest.GROUPS[:n])].sum() - t.loc["actual"]
        assert resid.abs().max() < 1e-4, (metric, resid)
    sol = _solve(compiled, "Base")
    bva = rows["BvA_Variance_Backtest"]
    bad = []
    for r in long.itertuples():
        col = openpyxl.utils.get_column_letter(cols[r.period_end])
        got = sol.get(("BVA_VARIANCE_BACKTEST", f"{col}{bva[f'br_{r.metric}_{r.component}']}"))
        if not isinstance(got, (int, float)) or abs(got - r.value) > 1e-6 * max(1.0, abs(r.value)):
            bad.append((r.metric, r.component, r.period_end, got, r.value))
    assert len(long) == 4 * (6 + 10 + 11 + 3 * 2) and not bad, bad[:5]


def test_workbook_labels_follow_the_config(built):
    _, wb, rows, _ = built
    om, ca = wb["Operating_Model"], wb["Cohort_ARR_Build"]
    assert om.cell(rows["Operating_Model"]["licenses"], 2).value == "Term license revenue"
    assert ca.cell(rows["Cohort_ARR_Build"]["arr"], 2).value == "ARR proxy (Platform annualised + Add-ons + Licenses trailing 4Q)"
    assert "platform_run_rate" in rows["Cohort_ARR_Build"] and "atlas" not in rows["Operating_Model"]
    text = " ".join(str(c.value) for ws in wb.worksheets if ws.title != "Sources" for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and not c.value.startswith("="))
    for word in ("MongoDB", "Atlas", "Enterprise Advanced", "FY27"):
        assert word not in text, word


def test_no_numeric_constants_in_formulas(built):
    _, wb, _, _ = built
    offenders = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    body = re.sub(r"('[^']+'|[A-Za-z_]+)?!?\$?[A-Z]{1,3}\$?\d+", "", c.value)
                    body = re.sub(r'"[^"]*"', "", body)
                    if [x for x in re.findall(r"\d+(?:\.\d+)?", body) if x not in ("0", "1")]:
                        offenders.append((ws.title, c.coordinate, c.value))
    assert not offenders, offenders[:5]


def test_without_guidance_the_guidance_rows_are_left_out(synth_db, tmp_path):
    """A company whose earnings releases give no parsed guidance: no guidance rows, no -100% errors."""
    src = synth.source_tables()
    src["guidance"] = src["guidance"].iloc[0:0]
    con = duckdb.connect()
    build.assemble(con, src)
    assert backtest.scorecard(con).guidance.isna().all()
    path = workbook.build(con, tmp_path / NAME)
    keys = {c.value for c in openpyxl.load_workbook(path)["BvA_Variance_Backtest"]["A"] if c.value}
    assert "sc_model" in keys and "sc_damped" in keys and not any("guid" in k for k in keys)


def test_cli_config_flag_selects_the_company(synth_db):
    from saas_metrics import cli
    assert config.DB_PATH == config.ROOT / "data" / "SYNTHCO" / "saasfpa.duckdb"   # per-company default paths
    assert config.WORKBOOK.name == "synthco_operating_model.xlsx"
    with pytest.raises(SystemExit, match="MongoDB write-up"):
        cli.main(["--config", str(synth.CONFIG), "memo"])
    with pytest.raises(SystemExit, match="no config"):
        cli.main(["--config", "NOPE", "qc"])
    assert config.TICKER == "SYNTHCO"
