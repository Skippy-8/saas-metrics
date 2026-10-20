"""The workbook recalculates to the engine (all checks PASS) with no typed numbers in formulas; the CLI runs."""
from __future__ import annotations

import re
from datetime import date

import openpyxl
import pytest

import saas_metrics as sm
from saas_metrics import cli, sample

formulas = pytest.importorskip("formulas")


@pytest.fixture(scope="module")
def rep():
    return sm.analyze(sample.generate(months=16), mapping=sample.MAPPING, today=date(2026, 10, 20))


def _solve(path):
    out = {}
    for k, v in formulas.ExcelModel().loads(str(path)).finish().calculate().items():
        m = re.match(r"'\[[^\]]+\]([^']+)'!([A-Z]+\d+)$", k)
        if m:
            out[(m.group(1), m.group(2))] = v.value[0][0] if hasattr(v, "value") else v
    return out


@pytest.mark.parametrize("max_customers", [50_000, 10])     # 10: retention written as values (large tapes)
def test_all_checks_pass(rep, tmp_path, max_customers):
    path = rep.to_excel(tmp_path / "m.xlsx", max_customers=max_customers)
    sol = _solve(path)
    wb = openpyxl.load_workbook(path)
    cell = wb.defined_names["AllChecks"].attr_text.split("!")[1].replace("$", "")
    assert sol[("CHECKS", cell)] == "PASS"
    nrr = [v for (s, a), v in sol.items() if s == "METRICS" and a.startswith("I") and isinstance(v, float)]
    assert nrr and all(0.5 < v < 2 for v in nrr)


def test_no_typed_numbers_in_formulas(rep, tmp_path):
    bad = []
    for ws in openpyxl.load_workbook(rep.to_excel(tmp_path / "m.xlsx")).worksheets:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str) and c.value.startswith("="):
                    body = re.sub(r'"[^"]*"', "", c.value)
                    body = re.sub(r"(\w+!)?\$?[A-Z]{1,3}\$?\d+", "", body)
                    if [x for x in re.findall(r"(?<![A-Za-z_])\d+(?:\.\d+)?", body) if x not in ("0", "1")]:
                        bad.append((ws.title, c.coordinate, c.value))
    assert not bad, bad[:3]


def test_cli_sample_and_report(tmp_path, capsys):
    f = tmp_path / "billing.csv"
    cli.main(["sample", str(f)])
    cli.main(["report", str(f), "--customer", "Customer", "--date", "Invoice Date", "--amount", "Amount",
              "--service-start", "Service Start", "--service-end", "Service End", "--segment", "Plan Tier",
              "--channel", "Channel", "--out", str(tmp_path / "out"), "--excel", str(tmp_path / "m.xlsx")])
    out = capsys.readouterr().out
    assert "identity holds in every period: True" in out and "nrr_12m" in out
    assert (tmp_path / "m.xlsx").exists() and (tmp_path / "out" / "metrics.csv").exists()
    with pytest.raises(SystemExit, match="input error"):
        cli.main(["report", str(f), "--customer", "Client"])
