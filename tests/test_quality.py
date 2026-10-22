"""Deferred revenue, RPO, SBC, implied billings and SBC-adjusted FCF — offline, on frozen data.

Two layers:

* synthetic companyfacts built period by period, so each tagging quirk is one small test: the
  legacy-to-ASC-606 switch, a half-tagged balance sheet, a filer that tags no deferred revenue at
  all, a restatement filed after the as-of date, a company-extension RPO tag;
* the four real filers in tests/fixtures/benchmark_*.json.gz, which is where the messy parts
  actually live (DDOG tagged both taxonomies in the transition quarter, TTD tags neither).

No network, no browser, no subprocess.
"""
from __future__ import annotations

import gzip
import json
from pathlib import Path

import pandas as pd
import pytest

from saas_metrics import benchmark, company, quality
from saas_metrics.company import Company
from saas_metrics.fiscal import FiscalCalendar

CAL = FiscalCalendar(12)                      # the synthetic filer has a December year end

FIX = Path(__file__).parent / "fixtures"
TICKERS = ["DDOG", "MDB", "SNOW", "TTD"]

REVENUE = "us-gaap:RevenueFromContractWithCustomerExcludingAssessedTax"
CWC_CUR = "us-gaap:ContractWithCustomerLiabilityCurrent"
CWC_NONCUR = "us-gaap:ContractWithCustomerLiabilityNoncurrent"
CWC = "us-gaap:ContractWithCustomerLiability"
DEF_CUR = "us-gaap:DeferredRevenueCurrent"
DEF_NONCUR = "us-gaap:DeferredRevenueNoncurrent"
RPO = "us-gaap:RevenueRemainingPerformanceObligation"
SBC = "us-gaap:ShareBasedCompensation"
SBC_ALLOC = "us-gaap:AllocatedShareBasedCompensationExpense"
CFO = "us-gaap:NetCashProvidedByUsedInOperatingActivities"
CAPEX = "us-gaap:PaymentsToAcquirePropertyPlantAndEquipment"
DELTA_DR = "us-gaap:IncreaseDecreaseInContractWithCustomerLiability"

LINES = {"revenue": REVENUE, "cost_of_revenue": "us-gaap:CostOfGoodsAndServicesSold",
         "rnd": "us-gaap:ResearchAndDevelopmentExpense", "snm": "us-gaap:SellingAndMarketingExpense",
         "gna": "us-gaap:GeneralAndAdministrativeExpense", "operating_income": "us-gaap:OperatingIncomeLoss",
         "sbc": SBC, "cfo": CFO, "capex": CAPEX}


# ---- synthetic companyfacts ---------------------------------------------------------------------
def quarters(n: int = 10, last: str = "2026-06-30") -> list[str]:
    ends = [last]
    while len(ends) < n:
        ends.append(CAL.prev_quarter_end(ends[-1]).isoformat())
    return sorted(ends)


def quarter_start(end: str) -> str:
    return CAL.quarter_start(end).isoformat()


class Filing:
    """Accumulate facts, one accession per period so every fact is inside its filing's scope."""

    def __init__(self) -> None:
        self.rows: list[dict] = []

    def duration(self, concept: str, start: str, end: str, value: float, *, filed: str | None = None,
                 accn: str | None = None) -> Filing:
        self.rows.append({"concept": concept, "start": start, "end": end, "val": value,
                          "accn": accn or f"acc-{end}", "filed": filed or "2026-08-01", "form": "10-Q"})
        return self

    def instant(self, concept: str, end: str, value: float, *, filed: str | None = None,
                accn: str | None = None) -> Filing:
        self.rows.append({"concept": concept, "end": end, "val": value,
                          "accn": accn or f"acc-{end}", "filed": filed or "2026-08-01", "form": "10-Q"})
        return self

    def drop(self, concept: str, *, except_for: tuple[str, ...] = (), only_for: tuple[str, ...] = ()) -> Filing:
        """`only_for` drops that concept for those periods; `except_for` drops it for every other period."""
        def keep(row: dict) -> bool:
            if row["concept"] != concept:
                return True
            if except_for:
                return row["end"] in except_for
            if only_for:
                return row["end"] not in only_for
            return False
        self.rows = [r for r in self.rows if keep(r)]
        return self

    def build(self) -> dict:
        facts: dict[str, dict] = {}
        for r in self.rows:
            ns, _, tag = r["concept"].partition(":")
            entry = facts.setdefault(ns, {}).setdefault(tag, {"label": tag, "units": {"USD": []}})
            entry["units"]["USD"].append({k: v for k, v in r.items() if k != "concept"})
        return {"cik": 1, "entityName": "Synthetic, Inc.", "facts": facts}


def baseline(ends: list[str]) -> tuple[Filing, dict[str, list[float]]]:
    """Ten quarters of a healthy subscription business with every tag the pipeline reads."""
    f = Filing()
    series = {"revenue": [], "dr_cur": [], "dr_noncur": [], "rpo": [], "sbc": []}
    for i, e in enumerate(ends):
        s = quarter_start(e)
        revenue = 100_000_000 + 10_000_000 * i
        cost = 0.4 * revenue
        rnd, snm, gna = 20_000_000, 20_000_000, 10_000_000
        sbc = 20_000_000 + 1_000_000 * i
        series["revenue"].append(revenue)
        series["dr_cur"].append(30_000_000 + 3_000_000 * i)
        series["dr_noncur"].append(5_000_000 + 500_000 * i)
        series["rpo"].append(300_000_000 + 20_000_000 * i)
        series["sbc"].append(sbc)
        f.duration(REVENUE, s, e, revenue)
        f.duration("us-gaap:CostOfGoodsAndServicesSold", s, e, cost)
        f.duration("us-gaap:ResearchAndDevelopmentExpense", s, e, rnd)
        f.duration("us-gaap:SellingAndMarketingExpense", s, e, snm)
        f.duration("us-gaap:GeneralAndAdministrativeExpense", s, e, gna)
        f.duration("us-gaap:OperatingIncomeLoss", s, e, revenue - cost - rnd - snm - gna)
        f.duration(SBC, s, e, sbc)
        f.duration(CFO, s, e, 50_000_000)
        f.duration(CAPEX, s, e, 5_000_000)
        if i:                                   # the cash-flow statement's own change line, for the cross-check
            f.duration(DELTA_DR, s, e, 3_500_000)
        f.instant(CWC_CUR, e, series["dr_cur"][-1])
        f.instant(CWC_NONCUR, e, series["dr_noncur"][-1])
        f.instant(RPO, e, series["rpo"][-1])
    return f, series


def synthetic_company(**quality_section) -> Company:
    raw = {"company": {"ticker": "TEST", "name": "Synthetic, Inc.", "cik": 1, "fiscal_year_end_month": 12},
           "benchmark": {"lines": dict(LINES)}}
    if quality_section:
        raw["benchmark"]["quality"] = quality_section
    return Company("TEST", "Synthetic, Inc.", 1, 12, raw)


def derive(f: Filing, **quality_section) -> quality.QualityResult:
    return quality.derive(synthetic_company(**quality_section), f.build())


# ---- the arithmetic -----------------------------------------------------------------------------
def test_implied_billings_and_sbc_adjusted_fcf_add_up():
    ends = quarters()
    f, s = baseline(ends)
    res = derive(f)
    row = res.row
    assert row["period_end"] == ends[-1]
    assert row["revenue_ttm"] == pytest.approx(sum(s["revenue"][-4:]))                 # 700M
    assert row["deferred_revenue"] == pytest.approx(s["dr_cur"][-1] + s["dr_noncur"][-1])
    # implied billings per quarter = revenue + the step in total deferred revenue, +3.5M here
    assert res.quarterly.at[ends[-1], "delta_deferred_revenue"] == pytest.approx(3_500_000)
    assert res.quarterly.at[ends[-1], "implied_billings"] == pytest.approx(s["revenue"][-1] + 3_500_000)
    assert row["implied_billings_ttm"] == pytest.approx(sum(s["revenue"][-4:]) + 4 * 3_500_000)
    assert row["billings_to_revenue"] == pytest.approx(row["implied_billings_ttm"] / row["revenue_ttm"])
    # FCF and the SBC charge that operating cash flow adds back
    assert row["fcf_ttm"] == pytest.approx(200_000_000 - 20_000_000)
    assert row["sbc_ttm"] == pytest.approx(sum(s["sbc"][-4:]))
    assert row["sbc_adjusted_fcf_ttm"] == pytest.approx(row["fcf_ttm"] - row["sbc_ttm"])
    assert row["rpo"] == pytest.approx(s["rpo"][-1])
    assert row["rpo_cover_years"] == pytest.approx(s["rpo"][-1] / row["revenue_ttm"])


def test_growth_rates_use_four_quarters_back_not_the_prior_quarter():
    ends = quarters()
    f, s = baseline(ends)
    row = derive(f).row
    assert row["deferred_revenue_yoy"] == pytest.approx(
        (s["dr_cur"][-1] + s["dr_noncur"][-1]) / (s["dr_cur"][-5] + s["dr_noncur"][-5]) - 1)
    assert row["rpo_growth_yoy"] == pytest.approx(s["rpo"][-1] / s["rpo"][-5] - 1)


def test_every_check_passes_on_a_clean_filer():
    res = derive(baseline(quarters())[0])
    failed = [c for c in res.checks if not c[1]]
    assert not failed, failed
    assert any("deferred revenue components sum to the total tag" in c[0] for c in res.checks) is False  # no total tag
    assert any("Δ deferred revenue agrees" in n for n in res.notes)


# ---- tag fallback: legacy deferred revenue -> ASC 606 --------------------------------------------
def test_pre_asc606_quarters_are_spliced_from_the_legacy_tags():
    ends = quarters()
    f, s = baseline(ends)
    early = tuple(ends[:4])
    f.drop(CWC_CUR, only_for=early).drop(CWC_NONCUR, only_for=early)
    for i, e in enumerate(early):
        f.instant(DEF_CUR, e, s["dr_cur"][i] - 2_000_000)
        f.instant(DEF_NONCUR, e, s["dr_noncur"][i])
    res = derive(f)
    dr = res.resolutions["deferred_revenue"]
    assert dr.primary is not None and "ContractWithCustomerLiability" in dr.primary.label
    assert dr.fallbacks == list(early)
    for i, e in enumerate(early):
        assert res.quarterly.at[e, "deferred_revenue"] == pytest.approx(s["dr_cur"][i] - 2_000_000
                                                                        + s["dr_noncur"][i])
        assert res.quarterly.at[e, "deferred_revenue_spliced"]
        assert "DeferredRevenue" in res.quarterly.at[e, "deferred_revenue_source"]
    assert not res.quarterly.at[ends[4], "deferred_revenue_spliced"]
    assert any("changed tags" in n and "spliced" in n for n in res.notes)
    # the spliced quarters are still complete, so the coverage gate passes
    assert all(c[1] for c in res.checks), res.checks


def test_a_half_tagged_balance_sheet_is_blank_not_a_partial_balance():
    ends = quarters()
    f, _ = baseline(ends)
    f.drop(CWC_NONCUR, only_for=(ends[7],))
    res = derive(f)
    assert pd.isna(res.quarterly.at[ends[7], "deferred_revenue"])
    assert pd.isna(res.quarterly.at[ends[7], "implied_billings"])
    coverage = [c for c in res.checks if "deferred revenue is tagged for every quarter" in c[0]]
    assert coverage and not coverage[0][1]
    assert ends[7] in coverage[0][2]


def test_a_total_tag_is_only_a_fallback_when_the_components_are_missing():
    """A single total must not be interleaved with a current + noncurrent series."""
    ends = quarters()
    f, s = baseline(ends)
    f.drop(CWC_CUR, only_for=(ends[2],)).drop(CWC_NONCUR, only_for=(ends[2],))
    f.instant(CWC, ends[2], s["dr_cur"][2] + s["dr_noncur"][2])
    res = derive(f)
    assert res.quarterly.at[ends[2], "deferred_revenue"] == pytest.approx(s["dr_cur"][2] + s["dr_noncur"][2])
    assert "total" in res.quarterly.at[ends[2], "deferred_revenue_source"]


def test_two_families_that_disagree_are_reported_and_not_averaged():
    ends = quarters()
    f, s = baseline(ends)
    f.instant(DEF_CUR, ends[5], s["dr_cur"][5] - 9_000_000)
    f.instant(DEF_NONCUR, ends[5], s["dr_noncur"][5])
    res = derive(f)
    dr = res.resolutions["deferred_revenue"]
    assert res.quarterly.at[ends[5], "deferred_revenue"] == pytest.approx(s["dr_cur"][5] + s["dr_noncur"][5])
    assert dr.conflicts and dr.conflicts[0]["period"] == ends[5]
    assert any("rather than averaged" in n for n in res.notes)


# ---- what is absent stays absent ----------------------------------------------------------------
def test_a_filer_that_tags_no_deferred_revenue_gets_blank_billings_not_zero():
    ends = quarters()
    f, _ = baseline(ends)
    f.drop(CWC_CUR).drop(CWC_NONCUR)
    res = derive(f)
    row = res.row
    assert res.quarterly["deferred_revenue"].isna().all()
    assert res.quarterly["implied_billings"].isna().all()
    assert row["deferred_revenue"] is None and row["implied_billings_ttm"] is None
    assert row["billings_to_revenue"] is None and row["deferred_revenue_yoy"] is None
    assert any("absent, not zero" in n for n in res.notes)
    # the rest of the row still works: SBC and FCF do not depend on deferred revenue
    assert row["sbc_adjusted_fcf_ttm"] == pytest.approx(row["fcf_ttm"] - row["sbc_ttm"])


def test_no_rpo_tag_means_blank_rpo_metrics():
    ends = quarters()
    f, _ = baseline(ends)
    f.drop(RPO)
    res = derive(f)
    assert res.row["rpo"] is None and res.row["rpo_growth_yoy"] is None and res.row["rpo_cover_years"] is None
    assert any("no RPO concept is tagged" in n for n in res.notes)


def test_sbc_falls_back_to_the_income_statement_tag_and_says_which_quarter():
    ends = quarters()
    f, s = baseline(ends)
    missing = ends[0]                                              # outside the trailing two years
    f.drop(SBC, only_for=(missing,))                               # the cash-flow add-back is missing there
    for e in ends[:6]:
        f.duration(SBC_ALLOC, quarter_start(e), e, s["sbc"][ends.index(e)])
    res = derive(f)
    sbc = res.resolutions["sbc"]
    assert sbc.primary is not None and sbc.primary.label.startswith("ShareBasedCompensation")
    assert res.quarterly.at[missing, "sbc_source"] == "AllocatedShareBasedCompensationExpense"
    assert res.quarterly.at[missing, "sbc"] == pytest.approx(s["sbc"][0])
    assert any("from the other tag" in n for n in res.notes)


# ---- point in time ------------------------------------------------------------------------------
def test_as_of_ignores_a_restatement_filed_later():
    ends = quarters()
    f, s = baseline(ends)
    f.instant(CWC_CUR, ends[-1], 999_000_000, filed="2026-09-01", accn="restated")
    original = quality.derive(synthetic_company(), f.build(), as_of="2026-08-15")
    restated = quality.derive(synthetic_company(), f.build())
    assert original.row["deferred_revenue"] == pytest.approx(s["dr_cur"][-1] + s["dr_noncur"][-1])
    assert restated.row["deferred_revenue"] == pytest.approx(999_000_000 + s["dr_noncur"][-1])


def test_a_company_extension_rpo_tag_can_be_configured():
    ends = quarters()
    f, _ = baseline(ends)
    f.drop(RPO)
    for i, e in enumerate(ends):
        f.instant("testco:ContractedBacklog", e, 111_000_000 + 1_000_000 * i)
    res = derive(f, rpo=[["testco:ContractedBacklog"]])
    assert res.row["rpo"] == pytest.approx(111_000_000 + 1_000_000 * 9)
    assert res.resolutions["rpo"].primary.concepts == ("testco:ContractedBacklog",)
    # an unconfigured extension is invisible, so the default metric stays blank
    assert derive(f).row["rpo"] is None


def test_config_rejects_an_unknown_quality_metric():
    raw = {"company": {"ticker": "T", "name": "T", "cik": 1, "fiscal_year_end_month": 12},
           "benchmark": {"lines": dict(LINES), "quality": {"billings": [["us-gaap:Revenue..."]]}}}
    with pytest.raises(company.ConfigError, match="benchmark.quality.billings"):
        company.validate(raw, "T.yaml")


# ---- the real filers -----------------------------------------------------------------------------
@pytest.fixture(scope="module")
def facts():
    return {t: json.load(gzip.open(FIX / f"benchmark_{t}.json.gz", "rt")) for t in TICKERS}


@pytest.fixture(scope="module")
def latest(facts):
    return quality.table(TICKERS, offline=facts)


def test_every_real_filer_passes_every_check(latest):
    _, results = latest
    failed = [(t, c) for t, r in results.items() for c in r.checks if not c[1]]
    assert not failed, failed


def test_the_real_table_is_the_latest_filed_quarter(latest):
    frame, results = latest
    assert list(frame.fiscal_label) == ["Q2 FY26", "Q2 FY27", "Q2 FY27", "Q2 FY26"]
    for t, r in results.items():
        filed = benchmark.company_comps(company.load(t), json.load(
            gzip.open(FIX / f"benchmark_{t}.json.gz", "rt")))
        assert r.row["period_end"] == filed.row["period_end"], t


def test_mdb_and_ddog_splice_the_legacy_tag_before_asc606_and_agree_at_the_crossover(latest):
    _, results = latest
    for t, first_legacy in (("MDB", "2017-01-31"), ("DDOG", "2018-12-31")):
        dr = results[t].resolutions["deferred_revenue"]
        assert dr.primary.regime == "asc606" and dr.primary.components, t
        assert dr.fallbacks and min(dr.fallbacks) <= first_legacy, t
        legacy = [p for p in dr.values.index if "DeferredRevenue" in str(dr.sources.get(p))]
        assert legacy, t
        # every quarter the table reports is on the primary family, so the splice only touches history
        q = results[t].quarterly
        assert q["deferred_revenue"].tail(8).notna().all(), t
        assert not q["deferred_revenue_spliced"].tail(8).any(), t
        assert q["implied_billings"].tail(8).notna().all(), t


def test_snow_tags_only_asc606_and_ttd_tags_neither(latest):
    _, results = latest
    snow = results["SNOW"].resolutions["deferred_revenue"]
    assert snow.primary.regime == "asc606" and not snow.fallbacks
    ttd = results["TTD"]
    assert ttd.resolutions["deferred_revenue"].primary is None
    assert ttd.resolutions["rpo"].primary is None
    assert ttd.row["implied_billings_ttm"] is None
    assert ttd.row["sbc_adjusted_fcf_ttm"] == pytest.approx(ttd.row["fcf_ttm"] - ttd.row["sbc_ttm"])


def test_billings_are_ahead_of_revenue_growth_for_the_asc606_filers(latest):
    """Implied billings lead revenue when the backlog builds — the reason to compute them at all."""
    _, results = latest
    for t in ("MDB", "DDOG", "SNOW"):
        row = results[t].row
        assert row["implied_billings_ttm"] > row["revenue_ttm"], t
        assert 1.0 < row["billings_to_revenue"] < 1.5, t
        assert row["sbc_pct_revenue"] > 0.05, t
        assert row["rpo"] > row["deferred_revenue"], t


def test_headline_numbers_are_stable(latest):
    """Pinned to the frozen fixtures, so a tag or fiscal-calendar change shows up as a diff."""
    frame, _ = latest
    row = frame.set_index("ticker").loc["MDB"]
    assert row["revenue_ttm"] == pytest.approx(2_782_770_000, rel=1e-6)
    assert row["deferred_revenue"] == pytest.approx(447_695_000, abs=1e3)
    assert row["implied_billings_ttm"] == pytest.approx(2_921_314_000, rel=1e-6)
    assert row["rpo"] == pytest.approx(1_519_200_000, abs=1e3)
    assert row["sbc_ttm"] == pytest.approx(564_392_000, abs=1e3)
    assert row["fcf_ttm"] == pytest.approx(659_020_000, abs=5e4)
    assert row["sbc_adjusted_fcf_ttm"] == pytest.approx(94_628_000, abs=5e4)
    assert row["sbc_adjusted_fcf_margin"] == pytest.approx(0.0340, abs=0.001)


# ---- outputs ------------------------------------------------------------------------------------
def test_csv_output_writes_the_summary_and_the_quarters(tmp_path):
    facts = {t: json.load(gzip.open(FIX / f"benchmark_{t}.json.gz", "rt")) for t in ("MDB", "TTD")}
    frame, results = quality.table(["MDB", "TTD"], offline=facts)
    out = quality.write(frame, results, tmp_path / "quality.csv")
    assert out.exists() and out.with_name("quality_quarterly.csv").exists()
    long = pd.read_csv(out.with_name("quality_quarterly.csv"))
    assert list(long["ticker"].unique()) == ["MDB", "TTD"]
    assert {"implied_billings", "deferred_revenue_source", "sbc_source"} <= set(long.columns)


def test_xlsx_output_has_live_formulas_and_a_quarterly_sheet(tmp_path):
    import openpyxl

    facts = {t: json.load(gzip.open(FIX / f"benchmark_{t}.json.gz", "rt")) for t in ("MDB", "DDOG")}
    frame, results = quality.table(["MDB", "DDOG"], offline=facts)
    out = quality.write(frame, results, tmp_path / "quality.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Revenue_Quality", "Quarterly"]
    ws = wb["Revenue_Quality"]
    header = [c.value for c in ws[3]]
    col = {h: i + 1 for i, h in enumerate(header)}
    cell = ws.cell(4, col["FCF margin (TTM)"])
    assert str(cell.value).startswith("=") and "!" not in str(cell.value)
    assert "Operating cash flow, TTM" in header
    assert wb["Quarterly"].max_row > 10
