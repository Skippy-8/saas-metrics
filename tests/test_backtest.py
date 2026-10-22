"""Budget-vs-actual bridge, forecast scorecard, unit economics and the GTM case.

Python reference identities first, then the workbook (recalculated with `formulas`) against them.
"""
from __future__ import annotations

import re

import numpy as np
import openpyxl
import pandas as pd
import pytest

from saas_metrics import backtest, config, drivers, model, roi, workbook

formulas = pytest.importorskip("formulas")
NAME = "saas_operating_model.xlsx"
FC_COLS = ("S", "T", "U", "V")   # the four forecast quarters (15 history columns from D)
TOTAL = "W"


@pytest.fixture(scope="module")
def bridge(db):
    return backtest.bridge(db)


@pytest.fixture(scope="module")
def wb_solution(db, tmp_path_factory):
    path = tmp_path_factory.mktemp("bva") / NAME
    workbook.build(db, path)
    sol = {}
    for k, v in formulas.ExcelModel().loads(str(path)).finish().calculate().items():
        m = re.match(r"'\[[^\]]+\]([^']+)'!([A-Z]+\d+)$", k)
        if m:
            sol[(m.group(1), m.group(2))] = v.value[0][0]
    wb = openpyxl.load_workbook(path)
    rows = {s: {wb[s].cell(r, 1).value: r for r in range(1, wb[s].max_row + 1) if wb[s].cell(r, 1).value}
            for s in ("BvA_Variance_Backtest", "Unit_Economics_ROI")}
    return sol, rows


def _cell(wb_solution, sheet, key, col):
    sol, rows = wb_solution
    return sol.get((sheet.upper(), f"{col}{rows[sheet][key]}"))


# ---- reference identities -------------------------------------------------------------------
def test_forecast_quarters_are_the_four_after_t0():
    assert [model.periods.label(e) for e in model.forecast_quarters()] == ["Q3 FY26", "Q4 FY26", "Q1 FY27", "Q2 FY27"]


@pytest.mark.parametrize("metric", list(backtest.EXPLAINS))
def test_bridge_reconciles_exactly(bridge, metric):
    long, _ = bridge
    t = long[long.metric == metric].pivot(index="component", columns="period_end", values="value")
    comps = list(backtest.GROUPS[:backtest.EXPLAINS[metric]])
    resid = t.loc["plan"] + t.loc[comps].sum() - t.loc["actual"]
    assert resid.abs().max() < 1e-4, resid


def test_bridge_plan_is_the_frozen_base_forecast(db, bridge):
    long, _ = bridge
    hist = drivers.history(db)
    cal = db.execute("SELECT * FROM cohort_calibration WHERE vintage='as_of_cutoff'").df().set_index("period_end")
    hc = drivers.headcount_history(db, hist).set_index("period_end")["headcount"]
    base = model.run(hist, cal, hc, drivers.derive(db), "Base")
    for metric in backtest.EXPLAINS:
        plan = long[(long.metric == metric) & (long.component == "plan")].set_index("period_end").value
        assert plan.values == pytest.approx(base.loc[plan.index, metric].values, rel=1e-12)


def test_actuals_are_the_latest_filings(db, bridge):
    long, _ = bridge
    act = long[(long.metric == "revenue") & (long.component == "actual")].set_index("period_end").value
    q = db.execute("SELECT period_end, value FROM quarterly_financials WHERE vintage='latest' AND line='revenue'").df()
    assert act.values == pytest.approx(q.set_index("period_end").value.reindex(act.index).values)


def test_scorecard_uses_only_pre_cutoff_guidance(db):
    sc = backtest.scorecard(db)
    assert sc.guidance.iloc[0] == pytest.approx((587e6 + 592e6) / 2)          # Q3 FY26, issued 2025-08-26
    assert sc.guidance.iloc[2:].isna().all()                                   # no FY27 guidance at the cutoff
    q1, q2 = db.execute("SELECT value FROM quarterly_financials WHERE vintage='as_of_cutoff' AND line='revenue' "
                        "AND period_end IN ('2025-04-30','2025-07-31') ORDER BY period_end").df().value
    assert sc.guidance.iloc[1] == pytest.approx(2.35e9 - q1 - q2 - sc.guidance.iloc[0])


def test_damped_trend_recursion():
    y = [100.0 + 5 * t for t in range(8)]                      # a straight line
    lv, tr, sse = backtest.damped_path(y, 0.5, 0.3, 1.0)       # no damping: tracks it exactly
    assert sse == pytest.approx(0) and lv[-1] == pytest.approx(y[-1]) and tr[-1] == pytest.approx(5)
    p, fc = backtest.damped_forecast(y, 4)
    assert p["phi"] == max(backtest.DAMPED_GRID["phi"])          # least damping fits a straight line best
    lv, tr, _ = backtest.damped_path(y, p["alpha"], p["beta"], p["phi"])
    assert fc == pytest.approx([lv[-1] + sum(p["phi"] ** k for k in range(1, h + 1)) * tr[-1] for h in (1, 2, 3, 4)])


def test_damped_trend_uses_only_history_to_t0(db):
    sc = backtest.scorecard(db)
    hist = drivers.history(db)
    cal = db.execute("SELECT * FROM cohort_calibration WHERE vintage='as_of_cutoff'").df().set_index("period_end")
    hc = drivers.headcount_history(db, hist).set_index("period_end")["headcount"]
    rows = model.run(hist, cal, hc, drivers.derive(db), "Base")
    rev = list(rows[rows.kind == "Actual"].revenue)
    assert rows[rows.kind == "Actual"].index[-1] == "2025-07-31" and len(rev) == 15
    assert list(sc.damped_trend) == pytest.approx(backtest.damped_forecast(rev, 4)[1])
    assert sc.attrs["damped"] == backtest.damped_fit(rev)


def test_gtm_case_at_company_efficiency_does_not_pay_back(db):
    inp = {i.key: i.value for i in roi.roi_inputs(db)}
    sol = roi.summary(roi.gtm_cases(inp), inp)
    assert sol.loc["A", "payback_month"] == 25 and sol.loc["B", "payback_month"] == 31
    act = backtest.actual_rows(db)
    nl = db.execute("SELECT period_end, new_logos FROM cohort_calibration WHERE vintage='as_of_cutoff'").df()
    ue = roi.scorecard(act, nl.set_index("period_end").new_logos, 0.035)
    att = roi.attainment_at_efficiency(inp, ue.at["2025-07-31", "magic_number"])
    assert roi.ae_efficiency({**inp, "roi_attain": att}) == pytest.approx(ue.at["2025-07-31", "magic_number"])
    low = {**inp, "roi_attain": att}
    assert pd.isna(roi.summary(roi.gtm_cases(low), low).loc["A", "payback_month"])


# ---- workbook vs reference ------------------------------------------------------------------
def test_workbook_bridge_matches_reference(bridge, wb_solution):
    long, _ = bridge
    fq = model.forecast_quarters()
    bad = []
    for r in long.itertuples():
        got = _cell(wb_solution, "BvA_Variance_Backtest", f"br_{r.metric}_{r.component}", FC_COLS[fq.index(r.period_end)])
        if not isinstance(got, (int, float)) or abs(got - r.value) > 1e-6 * max(1.0, abs(r.value)):
            bad.append((r.metric, r.component, r.period_end, got, r.value))
    assert len(long) == 120 and not bad, bad[:5]


@pytest.mark.parametrize("metric", list(backtest.EXPLAINS))
def test_workbook_bridge_check_rows_are_zero(wb_solution, metric):
    assert abs(_cell(wb_solution, "BvA_Variance_Backtest", f"br_{metric}_check", TOTAL)) < 1e-3


def test_workbook_scorecard_matches_reference(db, wb_solution):
    sc = backtest.scorecard(db)
    pairs = {"sc_actual": "actual_revenue", "sc_model": "model_base", "sc_seasonal": "seasonal_naive",
             "sc_trend": "trend_naive", "sc_guidance": "guidance", "ng_model": "model_base_ng_op",
             "ng_actual_model": "actual_ng_op_model_def", "ng_actual_company": "actual_ng_op_company_def",
             "ng_guidance": "guidance_ng_op", "sc_damped": "damped_trend"}
    for key, col in pairs.items():
        for c, e in zip(FC_COLS, sc.index, strict=True):
            exp = sc.at[e, col]
            got = _cell(wb_solution, "BvA_Variance_Backtest", key, c)
            if pd.isna(exp):
                assert got in (None, ""), (key, e, got)
            else:
                assert got == pytest.approx(exp, rel=1e-9), (key, e)


def test_workbook_unit_economics_matches_reference(db, wb_solution):
    hist = drivers.history(db)
    cal = db.execute("SELECT * FROM cohort_calibration WHERE vintage='as_of_cutoff'").df().set_index("period_end")
    hc = drivers.headcount_history(db, hist).set_index("period_end")["headcount"]
    drv = drivers.derive(db)
    rows = model.run(hist, cal, hc, drv, "Base")
    new_logos = cal.new_logos.reindex(rows.index)
    new_logos[rows.kind == "Forecast"] = next(d.base for d in drv if d.key == "new_logos")
    ue = roi.scorecard(rows, new_logos, 0.035)
    cols = [openpyxl.utils.get_column_letter(workbook.FIRST_COL + i) for i in range(len(rows))]
    pairs = {"ue_cac_payback": "cac_payback_months", "ue_magic": "magic_number", "ue_rule40": "rule_of_40",
             "ue_ltv_cac": "ltv_cac"}
    n = 0
    for key, col in pairs.items():
        for c, e in zip(cols, rows.index, strict=True):
            exp = ue.at[e, col]
            if pd.isna(exp):
                continue
            assert _cell(wb_solution, "Unit_Economics_ROI", key, c) == pytest.approx(exp, rel=1e-9), (key, e)
            n += 1
    assert n > 40


def test_workbook_gtm_case_matches_reference(db, wb_solution):
    inp = {i.key: i.value for i in roi.roi_inputs(db)}
    ref = roi.summary(roi.gtm_cases(inp), inp)
    for case in ("a", "b"):
        for col in ("cumulative_36m", "trough", "payback_month", "npv", "arr_month_36"):
            got = _cell(wb_solution, "Unit_Economics_ROI", f"{case}_{col}", "D")
            assert got == pytest.approx(ref.at[case.upper(), col], rel=1e-9), (case, col)
        irr = _cell(wb_solution, "Unit_Economics_ROI", f"{case}_irr_annual", "D")
        assert irr == pytest.approx(ref.at[case.upper(), "irr_annual"], rel=1e-6)
    assert _cell(wb_solution, "Unit_Economics_ROI", "a_magic", "D") == pytest.approx(roi.ae_efficiency(inp))
    co = _cell(wb_solution, "Unit_Economics_ROI", "co_magic", "D")
    assert _cell(wb_solution, "Unit_Economics_ROI", "a_attain_at_co_magic", "D") == pytest.approx(
        roi.attainment_at_efficiency(inp, co))


def test_workbook_checks_still_pass(wb_solution):
    sol, _ = wb_solution
    results = [v for (s, a), v in sol.items() if s == "CHECKS" and a.startswith("C") and a != "C3"]
    assert len(results) == 6 and set(results) == {"PASS"}


# ---- Phase 4: extensible point-in-time protocol ----------------------------------------------

class DummyForecaster:
    """A protocol-only model: constant revenue and margin, no operating-model internals."""

    def __init__(self):
        self.fits = 0
        self.history = None

    def fit(self, historical_df, as_of_date):
        self.fits += 1
        self.history = historical_df.copy()

    def predict(self, horizon_quarters):
        last = pd.to_datetime(self.history["period_end"]).max().date().isoformat()
        periods = model.quarters_after(last, horizon_quarters)
        return pd.DataFrame({"period_end": periods, "revenue": [100e6] * horizon_quarters,
                             "fcf_margin": [0.10] * horizon_quarters})


def _forecast(periods, revenue=100e6, margin=0.10):
    return pd.DataFrame({"period_end": periods, "revenue": [revenue] * len(periods),
                         "fcf_margin": [margin] * len(periods)})


def test_default_forecaster_satisfies_the_protocol():
    assert isinstance(backtest.DefaultSaaSForecaster(), backtest.BaseForecaster)
    assert not isinstance(object(), backtest.BaseForecaster)


def test_guard_rejects_a_future_period_end():
    frame = pd.DataFrame({"period_end": ["2025-01-31", "2025-10-31"], "revenue": [1.0, 2.0]})
    with pytest.raises(backtest.LookAheadBiasError) as exc:
        backtest.guard_no_future(frame, "2025-07-31")
    assert exc.value.offenders[0][0] == "period_end"
    assert "2025-10-31" in str(exc.value)
    assert "2025-10-31" in exc.value.offenders[0][1]


def test_guard_scans_filing_dates_and_a_datetime_index():
    frame = pd.DataFrame({"filing_date": ["2025-08-27"], "revenue": [1.0]})
    with pytest.raises(backtest.LookAheadBiasError):
        backtest.guard_no_future(frame, "2025-07-31")
    indexed = pd.DataFrame({"revenue": [1.0]}, index=pd.to_datetime(["2025-08-27"]))
    with pytest.raises(backtest.LookAheadBiasError):
        backtest.guard_no_future(indexed, "2025-07-31")


def test_guard_refuses_undated_history():
    with pytest.raises(backtest.ForecastContractError):
        backtest.guard_no_future(pd.DataFrame({"revenue": [1.0]}), "2025-07-31")


def test_dummy_forecaster_cannot_see_future_data(db):
    spy = DummyForecaster()

    def leaky(day):
        return pd.DataFrame({"period_end": ["2025-01-31", "2025-10-31"],
                             "revenue": [1.0, 2.0], "in_window": [True, True]})

    harness = backtest.PointInTimeBacktester(con=db, forecaster=spy, history_loader=leaky,
                                             max_evaluations=1)
    with pytest.raises(backtest.LookAheadBiasError):
        harness.run(dates=["2025-07-31"])
    assert spy.fits == 0, "fit must not run on a slice that failed the look-ahead guard"


def test_default_forecaster_fit_also_guards_future_data():
    frame = pd.DataFrame({"period_end": ["2025-10-31"], "revenue": [1.0]})
    with pytest.raises(backtest.LookAheadBiasError):
        backtest.DefaultSaaSForecaster().fit(frame, "2025-07-31")


def test_validate_forecast_derives_fcf_and_orders_periods():
    raw = _forecast(["2025-10-31", "2025-07-31"])
    out = backtest.validate_forecast(raw, 2, "2025-07-31",
                                     expected_periods=["2025-07-31", "2025-10-31"])
    assert out["period_end"].tolist() == ["2025-07-31", "2025-10-31"]
    assert np.allclose(out["fcf"], out["revenue"] * out["fcf_margin"])


def test_validate_forecast_rejects_contract_violations():
    with pytest.raises(backtest.ForecastContractError, match="missing required column"):
        backtest.validate_forecast(pd.DataFrame({"period_end": ["2025-10-31"], "revenue": [1.0]}),
                                   1, "2025-07-31")
    with pytest.raises(backtest.ForecastContractError, match="row"):
        backtest.validate_forecast(_forecast(["2025-10-31"]), 2, "2025-07-31")
    with pytest.raises(backtest.ForecastContractError, match="duplicate"):
        backtest.validate_forecast(_forecast(["2025-10-31", "2025-10-31"]), 2, "2025-07-31")
    with pytest.raises(backtest.ForecastContractError, match="do not match"):
        backtest.validate_forecast(_forecast(["2025-10-31", "2026-01-31"]), 2, "2025-07-31",
                                   expected_periods=["2025-07-31", "2025-10-31"])


def test_trailing_plan_matches_frozen_drivers_at_cutoff(db):
    history = backtest.sec_history(db, config.CUTOFF_DATE)
    plan = backtest.derive_plan(history)
    frozen = {d.key: d.base for d in drivers.derive(db)}
    assert set(plan) == set(frozen)
    for key, value in frozen.items():
        assert plan[key] == pytest.approx(value, rel=0, abs=1e-9), key


def test_default_forecaster_reproduces_the_frozen_model(db):
    history = backtest.sec_history(db, config.CUTOFF_DATE)
    cal = db.execute("SELECT * FROM cohort_calibration WHERE vintage='as_of_cutoff'").df() \
        .set_index("period_end")
    raw = drivers.history(db)
    hc = drivers.headcount_history(db, raw).set_index("period_end")["headcount"]
    drv = drivers.derive(db)
    for scenario in drivers.SCENARIOS:
        forecaster = backtest.DefaultSaaSForecaster(scenario=scenario)
        forecaster.fit(history, config.CUTOFF_DATE)
        forecast = forecaster.predict(4)
        ref = model.run(raw, cal, hc, drv, scenario)
        ref = ref[ref.kind == "Forecast"]
        assert np.allclose(forecast["revenue"], ref["revenue"]), scenario
        assert np.allclose(forecast["fcf"], ref["fcf"]), scenario


def test_any_protocol_forecaster_runs(db):
    result = backtest.PointInTimeBacktester(con=db, forecaster=DummyForecaster(),
                                            max_evaluations=1).run()
    assert result.evaluations["as_of"].nunique() == 1
    assert len(result.evaluations) == 8              # 4 quarters x 2 metrics
    assert set(result.summary["metric"]) == {"revenue", "fcf"}


def test_backtester_scores_multiple_origins(db):
    result = backtest.PointInTimeBacktester(con=db, max_evaluations=2).run()
    assert result.evaluations["as_of"].nunique() == 2
    assert len(result.evaluations) == 16             # 2 origins x 4 quarters x 2 metrics
    assert set(result.summary["metric"]) == {"revenue", "fcf"}
    assert (result.summary["mape"].dropna() >= 0).all()
    assert set(result.bridges["window"]) == {"1Q", "4Q"}
    assert np.allclose(result.bridges["variance"],
                       result.bridges["forecast"] - result.bridges["actual"])
    four = result.bridges[result.bridges["window"] == "4Q"].iloc[0]
    part = result.evaluations[(result.evaluations["as_of"] == four["as_of"])
                              & (result.evaluations["metric"] == four["metric"])]
    assert four["forecast"] == pytest.approx(part["forecast"].sum())
    assert four["actual"] == pytest.approx(part["actual"].sum())
    assert "MAPE" in backtest.format_result(result)


def test_billing_slice_uses_period_end():
    frame = pd.DataFrame({"period_end": ["2025-01-31", "2025-07-31", "2025-10-31"],
                          "amount": [1.0, 2.0, 3.0]})
    out = backtest.slice_billing(frame, "2025-07-31")
    assert out["period_end"].tolist() == ["2025-01-31", "2025-07-31"]
    loader = backtest.billing_history_loader(frame)
    assert loader("2025-07-31")["period_end"].tolist() == ["2025-01-31", "2025-07-31"]
    with pytest.raises(backtest.ForecastContractError, match="period_end"):
        backtest.slice_billing(frame.drop(columns=["period_end"]), "2025-07-31")
    broken = frame.copy()
    broken.loc[0, "period_end"] = "not-a-date"
    with pytest.raises(backtest.ForecastContractError, match="unparseable"):
        backtest.slice_billing(broken, "2025-07-31")


def test_legacy_backtest_exports_are_unchanged():
    assert backtest.GROUPS[:3] == ("expansion", "new_logo_volume", "new_logo_arr")
    assert backtest.EXPLAINS["fcf"] == 10
    assert backtest.GROUP_LABELS["gross_margin"] == "Gross margin rate"
    assert callable(backtest.bridge) and callable(backtest.scorecard)
    assert callable(backtest.damped_forecast) and backtest.DAMPED_GRID["phi"]


def test_pit_cli_flag_keeps_the_default_path_untouched():
    from saas_metrics import cli
    default = cli.parser().parse_args(["backtest"])
    assert default.pit is False and default.horizon == 4 and default.evaluations == 4
    assert default.fn is cli.cmd_backtest
    opts = cli.parser().parse_args(["backtest", "--pit", "--horizon", "3", "--evaluations", "2"])
    assert opts.pit is True and opts.horizon == 3 and opts.evaluations == 2


def test_sec_history_is_a_strict_filed_by_slice(db):
    history = backtest.sec_history(db, "2025-07-31")
    assert (pd.to_datetime(history["filed"]) <= pd.Timestamp("2025-07-31")).all()
    # Q2 FY26 ended on 2025-07-31 but was not filed until 2025-08-27: it must not be in the slice.
    assert "2025-07-31" not in set(history["period_end"])
    assert history["period_end"].max() == "2025-04-30"
