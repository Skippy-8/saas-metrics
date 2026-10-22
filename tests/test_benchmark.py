"""Peer benchmarking on frozen companyfacts fixtures (no network)."""
from __future__ import annotations

import gzip
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from saas_metrics import benchmark, company
from saas_metrics.fiscal import FiscalCalendar

FIX = Path(__file__).parent / "fixtures"
TICKERS = ["MDB", "DDOG", "SNOW", "TTD"]


@pytest.fixture(scope="module")
def facts():
    return {t: json.load(gzip.open(FIX / f"benchmark_{t}.json.gz", "rt")) for t in TICKERS}


@pytest.fixture(scope="module")
def latest(facts):
    return benchmark.comps(TICKERS, offline=facts)


def test_every_check_passes_and_rows_are_current(latest):
    table, results = latest
    failed = [(t, c) for t, r in results.items() for c in r.checks if not c[1]]
    assert not failed, failed
    assert list(table.fiscal_label) == ["Q2 FY27", "Q2 FY26", "Q2 FY27", "Q2 FY26"]
    assert all(any("latest filed quarter" in c[0] for c in r.checks) for r in results.values())


def test_mdb_quarters_match_the_verified_store(db, facts):
    """The XBRL normalizer reproduces the MongoDB build (filing instances) quarter by quarter."""
    res = benchmark.company_comps(company.load("MDB"), facts["MDB"])
    store = db.execute("SELECT * FROM q_wide WHERE vintage='latest' AND in_window").df().set_index("period_end")
    pairs = {"revenue": "revenue", "cost_of_revenue": "cost_of_revenue", "rnd": "rnd", "snm": "snm", "gna": "gna",
             "operating_income": "operating_income", "sbc": "sbc_total", "amortization": "amort_intangibles",
             "cfo": "cfo", "capex": "capex"}
    q = res.quarterly.loc[store.index]
    for mine, theirs in pairs.items():
        assert q[mine].values == pytest.approx(store[theirs].values, abs=0.5), mine
    assert res.ttm["revenue"] == pytest.approx(store["revenue"].iloc[-4:].sum())


def test_as_of_uses_only_filings_known_then(facts):
    table, results = benchmark.comps(TICKERS, as_of="2026-03-31", offline=facts)
    assert list(table.fiscal_label) == ["Q4 FY26", "Q4 FY25", "Q4 FY26", "Q4 FY25"]
    for t, r in results.items():
        assert all(c[1] for c in r.checks), (t, r.checks)
    # nothing filed after the as-of date reached the numbers: rebuild from facts truncated by filing date
    for t in TICKERS:
        cut = json.loads(json.dumps(facts[t]))
        for d in cut["facts"]["us-gaap"].values():
            for u in d["units"]:
                d["units"][u] = [f for f in d["units"][u] if f["filed"] <= "2026-03-31"]
        alone = benchmark.company_comps(company.load(t), cut, None)
        for k, v in alone.row.items():
            assert v == (pytest.approx(results[t].row[k]) if isinstance(v, float) else results[t].row[k]), (t, k)


def test_summed_line_gap_is_zero_with_a_note(latest):
    _, results = latest
    snow = results["SNOW"]
    assert any("PaymentsToDevelopSoftware not reported" in n and "last reported value $0" in n for n in snow.notes)


def test_ratios_follow_from_ttm(latest):
    table, results = latest
    for r in results.values():
        x, row = r.ttm, r.row
        assert row["gross_margin"] == pytest.approx(1 - x["cost_of_revenue"] / x["revenue"])
        assert row["fcf_margin"] == pytest.approx((x["cfo"] - x["capex"]) / x["revenue"])
        assert row["rule_of_40"] == pytest.approx(row["revenue_growth_yoy"] + row["fcf_margin"])
        assert row["non_gaap_op_margin_approx"] == pytest.approx(
            (x["operating_income"] + x["sbc"] + x["amortization"]) / x["revenue"])


def test_xlsx_formulas_recalculate_to_the_table(latest, tmp_path):
    formulas = pytest.importorskip("formulas")
    table, results = latest
    path = benchmark.write(table, results, tmp_path / "comps.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Peer_Comps"]
    head = {ws.cell(3, c).value: c for c in range(1, ws.max_column + 1)}
    for r in range(4, 4 + len(results)):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                lits = re.findall(r"\d+(?:\.\d+)?", re.sub(r"\$?[A-Z]{1,3}\$?\d+", "", v))
                assert set(lits) <= {"0", "1"}, v
    sol = {}
    for k, v in formulas.ExcelModel().loads(str(path)).finish().calculate().items():
        m = re.match(r"'\[[^\]]+\]([^']+)'!([A-Z]+\d+)$", k)
        if m:
            sol[m.group(2)] = v.value[0][0]
    for i, (t, res) in enumerate(results.items()):
        for col, label in benchmark.COLUMNS.items():
            if col in benchmark.FORMULAS:
                letter = openpyxl.utils.get_column_letter(head[label])
                assert sol[f"{letter}{4 + i}"] == pytest.approx(res.row[col], rel=1e-12), (t, col)


def test_fiscal_labels_use_each_company_calendar(latest):
    table, _ = latest
    for row in table.itertuples():
        cal = FiscalCalendar(company.load(row.ticker).fiscal_year_end_month)
        assert cal.label(row.period_end) == row.fiscal_label


def test_missing_required_concept_fails_loudly(facts):
    broken = json.loads(json.dumps(facts["DDOG"]))
    del broken["facts"]["us-gaap"]["SellingAndMarketingExpense"]
    with pytest.raises(ValueError, match="required line snm"):
        benchmark.company_comps(company.load("DDOG"), broken)


def test_csv_and_parquet_outputs(latest, tmp_path):
    table, results = latest
    for ext in ("csv", "parquet"):
        p = benchmark.write(table, results, tmp_path / f"comps.{ext}")
        back = pd.read_csv(p) if ext == "csv" else __import__("duckdb").connect().execute(f"SELECT * FROM '{p}'").df()
        assert list(back.ticker) == TICKERS
