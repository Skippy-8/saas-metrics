"""The MCP server: four tools, a registered schema for each, and an offline ingest -> report round trip.

`fastmcp` is an optional extra, so the module skips without it and the base install stays green.
Nothing here touches the network, a browser or a subprocess: SEC data comes from the frozen companyfacts
fixtures and the billing export is a synthetic Stripe CSV written under `tmp_path`.
"""
from __future__ import annotations

import asyncio
import csv
import gzip
import json
from datetime import date
from pathlib import Path

import duckdb
import pandas as pd
import pytest

from saas_metrics.contracts import ContractError

pytest.importorskip("fastmcp", reason="the MCP server needs the optional mcp extra")
srv = pytest.importorskip("saas_metrics.mcp_server")
ToolError = pytest.importorskip("fastmcp.exceptions").ToolError

FIX = Path(__file__).parent / "fixtures"
TICKERS = ["DDOG", "MDB", "SNOW", "TTD"]


# ---- fixtures and helpers ------------------------------------------------------------------------
@pytest.fixture(scope="module")
def facts() -> dict[str, dict]:
    out = {}
    for t in TICKERS:
        with gzip.open(FIX / f"benchmark_{t}.json.gz", "rt") as f:
            out[t] = json.load(f)
    return out


def epoch(y: int, m: int, d: int = 1) -> int:
    return int(pd.Timestamp(date(y, m, d)).timestamp())


def stripe_export(path: Path, months: int = 14, customers: dict[str, int] | None = None) -> Path:
    """A synthetic Stripe invoice CSV: minor units, one invoice per customer-month."""
    customers = customers or {"cus_A": 250_000, "cus_B": 100_000, "cus_C": 50_000}
    starts = list(pd.date_range("2025-01-01", periods=months + 1, freq="MS"))
    rows = []
    for i in range(months):
        start, end = starts[i], starts[i + 1]
        for cust, cents in customers.items():
            rows.append({"id": f"in_{cust}_{i}", "customer": cust, "created": int(start.timestamp()),
                         "period_start": int(start.timestamp()), "period_end": int(end.timestamp()),
                         "subtotal": cents, "currency": "usd", "plan": "Pro",
                         "subscription": f"sub_{cust}"})
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return path


def ingest(export: Path, db: Path) -> str:
    return srv.ingest_billing_ledger(str(export), "stripe", str(db))


# ---- tool registration ---------------------------------------------------------------------------
def test_the_four_tools_register_with_their_contracts() -> None:
    tools = {t.name: t for t in asyncio.run(srv.mcp.list_tools())}
    assert set(tools) == {"ingest_billing_ledger", "generate_cohort_report",
                          "run_sec_peer_benchmark", "analyze_revenue_quality"}

    ing = tools["ingest_billing_ledger"].parameters
    assert ing["required"] == ["export_path"]
    assert ing["properties"]["adapter"] == {
        "default": "stripe", "enum": ["stripe", "chargebee", "salesforce"], "type": "string",
        "description": "which system produced it.",
    }

    rep = tools["generate_cohort_report"].parameters
    assert rep.get("required", []) == []
    assert rep["properties"]["grain"]["enum"] == ["month", "quarter"]

    bench = tools["run_sec_peer_benchmark"].parameters
    assert bench["required"] == ["tickers"]
    assert bench["properties"]["tickers"]["items"] == {"type": "string"}
    assert bench["properties"]["quarters_back"]["default"] == 4

    qual = tools["analyze_revenue_quality"].parameters
    assert qual["required"] == ["tickers"]
    assert qual["properties"]["tickers"]["type"] == "array"

    # the docstrings are the descriptions an agent actually reads
    assert "DuckDB" in tools["ingest_billing_ledger"].description
    assert "waterfall" in tools["generate_cohort_report"].description
    assert "Rule of 40" in tools["run_sec_peer_benchmark"].description
    assert "deferred revenue" in tools["analyze_revenue_quality"].description


# ---- ingest --------------------------------------------------------------------------------------
def test_ingest_writes_the_canonical_ledger_and_reports_total_mrr(tmp_path: Path) -> None:
    export, db = stripe_export(tmp_path / "invoices.csv"), tmp_path / "billing.duckdb"
    text = ingest(export, db)

    assert "Ingested 42 row(s)" in text
    assert "total MRR $56,000.00" in text
    assert "Write-back verified against the in-memory ledger." in text
    assert "Validation warnings: none." in text

    con = duckdb.connect(str(db), read_only=True)
    try:
        assert con.execute("SELECT count(*) FROM billing_ledger").fetchone()[0] == 42
        assert con.execute("SELECT sum(mrr_amount) FROM billing_ledger").fetchone()[0] == 56_000
        assert con.execute("SELECT typeof(mrr_amount) FROM billing_ledger LIMIT 1").fetchone()[0] == \
            "DECIMAL(18,2)"
        assert con.execute("SELECT count(DISTINCT customer_id) FROM billing_ledger").fetchone()[0] == 3
    finally:
        con.close()


def test_ingest_surfaces_warnings_without_blocking_the_write(tmp_path: Path) -> None:
    export = stripe_export(tmp_path / "invoices.csv", months=1,
                           customers={"cus_A": 4_900, "cus_B": 1_000})
    db = tmp_path / "billing.duckdb"
    lines = export.read_text().splitlines()
    lines[1] = lines[1].replace(",sub_cus_A", ",")      # blank the first data row's subscription id
    export.write_text("\n".join(lines) + "\n")

    text = ingest(export, db)
    assert "Validation warnings (1)" in text
    assert "missing_contract_id" in text
    assert db.exists()


def test_ingest_blocks_on_an_error_and_writes_nothing(tmp_path: Path) -> None:
    export, db = stripe_export(tmp_path / "bad.csv", months=1), tmp_path / "billing.duckdb"
    rows = export.read_text().splitlines()
    # end the service period before it starts: the term cannot be negative
    bad = rows[1].split(",")
    bad[4], bad[3] = bad[3], bad[4]
    export.write_text("\n".join([rows[0], ",".join(bad)]) + "\n")

    with pytest.raises(ContractError, match="negative_duration"):
        ingest(export, db)
    assert not db.exists()


def test_ingest_requires_an_existing_export(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="export not found"):
        srv.ingest_billing_ledger(str(tmp_path / "nope.csv"), "stripe", str(tmp_path / "b.duckdb"))


# ---- report --------------------------------------------------------------------------------------
def test_report_reads_the_ledger_back_and_writes_an_auditable_workbook(tmp_path: Path) -> None:
    export, db = stripe_export(tmp_path / "invoices.csv"), tmp_path / "billing.duckdb"
    ingest(export, db)
    export.unlink()                                     # the report must come from DuckDB, not the CSV
    out = tmp_path / "reports" / "cohort.xlsx"

    text = srv.generate_cohort_report(str(out), db_path=str(db))

    assert str(out.resolve()) in text
    assert "- Starting ARR: $48K (beginning of 2026-02-28)" in text
    assert "- Ending ARR: $48K (end of 2026-02-28)" in text
    assert "- GRR (12m): 100.0%" in text
    assert "- NRR (12m): 100.0%" in text
    assert "14 month(s)" in text
    assert out.exists()
    from openpyxl import load_workbook
    assert {"Waterfall", "Metrics", "Checks"} <= set(load_workbook(out).sheetnames)


def test_report_defaults_to_a_file_next_to_the_ledger(tmp_path: Path, monkeypatch) -> None:
    export, db = stripe_export(tmp_path / "invoices.csv"), tmp_path / "store" / "billing.duckdb"
    ingest(export, db)
    elsewhere = tmp_path / "cwd"
    elsewhere.mkdir()
    monkeypatch.chdir(elsewhere)

    text = srv.generate_cohort_report(db_path=str(db))

    assert (db.parent / "cohort_report.xlsx").exists()
    assert str((db.parent / "cohort_report.xlsx").resolve()) in text


def test_the_ledger_database_can_come_from_the_environment(tmp_path: Path, monkeypatch) -> None:
    export, db = stripe_export(tmp_path / "invoices.csv"), tmp_path / "env.duckdb"
    monkeypatch.setenv(srv.DB_ENV, str(db))
    assert srv.ledger_db_path() == db.resolve()
    text = srv.ingest_billing_ledger(str(export), "stripe")
    assert str(db.resolve()) in text
    assert db.exists()


def test_report_without_a_ledger_explains_the_next_step(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="ingest_billing_ledger first"):
        srv.generate_cohort_report(db_path=str(tmp_path / "missing.duckdb"))

    db = tmp_path / "other.duckdb"
    duckdb.connect(str(db)).execute("CREATE TABLE unrelated (x INTEGER)").close()
    with pytest.raises(ValueError, match="has no 'billing_ledger' table"):
        srv.generate_cohort_report(db_path=str(db))


# ---- benchmark -----------------------------------------------------------------------------------
def test_benchmark_markdown_comes_from_the_frozen_filings(monkeypatch, facts) -> None:
    real = srv.benchmark.comps
    monkeypatch.setattr(srv.benchmark, "comps",
                        lambda tickers, **kw: real(tickers, offline=facts, **kw))

    text = srv.run_sec_peer_benchmark(["ddog", "SNOW"], quarters_back=2)

    assert "## SEC XBRL peer comps — trailing four fiscal quarters" in text
    assert "| DDOG | Q2 FY26 | $3,967M |" in text
    assert "58.6%" in text                                 # DDOG Rule of 40
    assert "## 2-quarter trend" in text
    assert "| DDOG | 2026-06-30 |" in text
    assert "6 passed, 0 failed." in text


def test_benchmark_can_be_cut_off_at_an_as_of_date(monkeypatch, facts) -> None:
    real = srv.benchmark.comps
    monkeypatch.setattr(srv.benchmark, "comps",
                        lambda tickers, **kw: real(tickers, offline=facts, **kw))

    text = srv.run_sec_peer_benchmark(["MDB"], quarters_back=1, as_of="2026-01-01")

    assert "as known on 2026-01-01" in text
    assert "| MDB |" in text


def test_benchmark_rejects_an_empty_list_and_a_zero_window() -> None:
    with pytest.raises(ValueError, match="must not be empty"):
        srv.run_sec_peer_benchmark([])
    with pytest.raises(ValueError, match="at least 1"):
        srv.run_sec_peer_benchmark(["DDOG"], quarters_back=0)


# ---- quality -------------------------------------------------------------------------------------
def test_quality_markdown_keeps_absent_tags_blank(monkeypatch, facts) -> None:
    real = srv.quality.table
    monkeypatch.setattr(srv.quality, "table",
                        lambda tickers, **kw: real(tickers, offline=facts, **kw))

    text = srv.analyze_revenue_quality(["MDB", "TTD"])

    assert "## Revenue-quality scorecard — trailing four fiscal quarters" in text
    assert "| MDB | Q2 FY27 | $2,783M | $448M | $2,921M | 105.0% |" in text
    assert "20.3%" in text and "3.4%" in text
    ttd = next(line for line in text.splitlines() if line.startswith("| TTD |"))
    assert ttd.count("n/a") >= 3                          # absent tags stay blank, never zero
    assert "no deferred-revenue concept is tagged" in text
    assert "7 passed, 0 failed." in text


# ---- the protocol layer and the CLI --------------------------------------------------------------
def test_a_client_call_uses_the_registered_schema(monkeypatch, facts) -> None:
    real = srv.benchmark.comps
    monkeypatch.setattr(srv.benchmark, "comps",
                        lambda tickers, **kw: real(tickers, offline=facts, **kw))

    async def call():
        return await srv.mcp.call_tool("run_sec_peer_benchmark", {"tickers": ["DDOG"], "quarters_back": 1})

    result = asyncio.run(call())
    assert not result.is_error
    assert "| DDOG | Q2 FY26 |" in result.content[0].text


def test_a_blocked_tool_surfaces_an_actionable_error_to_the_client(tmp_path: Path) -> None:
    async def call():
        return await srv.mcp.call_tool("generate_cohort_report",
                                       {"db_path": str(tmp_path / "missing.duckdb")})

    with pytest.raises(ToolError, match="ingest_billing_ledger first"):
        asyncio.run(call())


def test_the_mcp_subcommand_is_wired_to_the_server(monkeypatch) -> None:
    from saas_metrics import cli

    a = cli.parser().parse_args(["mcp"])
    assert a.fn is cli.cmd_mcp
    called = []
    monkeypatch.setattr(srv, "main", lambda: called.append(True))
    cli.cmd_mcp(a)
    assert called == [True]
