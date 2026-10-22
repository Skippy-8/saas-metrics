"""Public peer comp table from SEC XBRL (companyfacts), one row per company.

For each company in configs/: take the configured us-gaap concepts, keep only 10-Q/10-K facts a
filing can speak for (its own period and standard comparatives), pick each fact's value as known
on the as-of date (latest filing on or before it), derive quarterly values (reported 3-month, YTD
differencing, or full year minus nine months) on the company's own fiscal calendar, and sum the
trailing four quarters to its latest quarter.

Non-GAAP figures are not in XBRL. `non_gaap_op_margin_approx` adds back SBC and amortization only;
company definitions exclude more, so it is an approximation and is labelled as one.
"""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import pandas as pd

from saas_metrics import company as companies
from saas_metrics.fiscal import FiscalCalendar

RAW = companies.ROOT / "data" / "raw" / "benchmark"
SCOPE_MONTHS = {"10-Q": 12, "10-K": 24}
FORMS = ("10-Q", "10-K")
TTM_LINES = ("revenue", "cost_of_revenue", "rnd", "snm", "gna", "operating_income", "sbc", "amortization",
             "cfo", "capex")

COLUMNS = {  # output column -> description (also the CSV/xlsx header order)
    "ticker": "Ticker", "name": "Company", "fiscal_label": "Latest fiscal quarter", "period_end": "Quarter end",
    "revenue_ttm": "Revenue, TTM", "revenue_growth_yoy": "Revenue growth (TTM vs prior TTM)",
    "gross_margin": "Gross margin", "rnd_pct": "R&D % revenue", "snm_pct": "S&M % revenue",
    "gna_pct": "G&A % revenue", "sbc_pct": "SBC % revenue", "gaap_op_margin": "GAAP operating margin",
    "non_gaap_op_margin_approx": "Non-GAAP op. margin (approx.: + SBC + amortization)",
    "fcf_margin": "FCF margin (CFO - capex - capitalized software)", "rule_of_40": "Rule of 40",
}


@dataclass
class Result:
    row: dict
    ttm: dict              # line -> TTM value (and *_prior for revenue)
    quarterly: pd.DataFrame
    checks: list[tuple[str, bool, str]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


# ---- data ---------------------------------------------------------------------------------------
def companyfacts_path(c: companies.Company, raw_dir: Path = RAW) -> Path:
    return raw_dir / f"{c.ticker}.json"


def fetch(c: companies.Company, refresh: bool = False, raw_dir: Path = RAW) -> Path:
    from saas_metrics.sec_client import fetch as http

    path = companyfacts_path(c, raw_dir)
    if refresh or not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(http(f"https://data.sec.gov/api/xbrl/companyfacts/CIK{c.cik10}.json"))
    return path


def load_facts(cf: dict, concepts: set[str]) -> pd.DataFrame:
    """Facts from 10-Q/10-K filings for the given concepts, with each filing's own period end.

    Concepts may live in any namespace (a filer's RPO is sometimes a company extension, not
    `us-gaap:RevenueRemainingPerformanceObligation`), so the lookup goes through the prefix rather
    than assuming `us-gaap`.
    """
    wanted: dict[str, set[str]] = {}
    for concept in concepts:
        ns, _, tag = concept.partition(":")
        wanted.setdefault(ns, set()).add(tag)
    rows, report_end = [], {}
    for ns, tags in wanted.items():
        for tag, d in cf["facts"].get(ns, {}).items():
            if tag not in tags:
                continue
            for unit_facts in d["units"].values():
                for f in unit_facts:
                    if f.get("form") not in FORMS:
                        continue
                    # a filing's own period is the latest period end it reports
                    report_end[f["accn"]] = max(report_end.get(f["accn"], ""), f["end"])
                    rows.append({"concept": f"{ns}:{tag}", "start": f.get("start"), "end": f["end"],
                                 "value": float(f["val"]), "form": f["form"], "filed": f["filed"], "accn": f["accn"]})
    df = pd.DataFrame(rows, columns=["concept", "start", "end", "value", "form", "filed", "accn"])
    if df.empty:
        return df.assign(report_end=pd.Series(dtype=str), in_scope=pd.Series(dtype=bool))
    df["report_end"] = df["accn"].map(report_end)
    back = (pd.to_datetime(df["report_end"]).dt.to_period("M") - pd.to_datetime(df["end"]).dt.to_period("M")).map(
        lambda x: x.n)
    df["in_scope"] = back <= df["form"].map(SCOPE_MONTHS)
    return df


def known(facts: pd.DataFrame, as_of: str | None) -> pd.DataFrame:
    """Each (concept, start, end) as known on `as_of`: the latest in-scope filing on or before it."""
    f = facts[facts["in_scope"]]
    if as_of:
        f = f[f["filed"] <= as_of]
    return f.sort_values("filed").drop_duplicates(["concept", "start", "end"], keep="last")


def quarterly(kn: pd.DataFrame, concept: str, cal: FiscalCalendar) -> pd.Series:
    """Quarterly values of one duration concept, derived as in the MongoDB build."""
    k = kn[(kn["concept"] == concept) & kn["start"].notna()]
    idx = {(r.start, r.end): r.value for r in k.itertuples()}
    ends = sorted({e for e in k["end"] if cal.is_quarter_end(e)})
    out = {}
    for e in ends:
        qs, fys = cal.quarter_start(e).isoformat(), cal.fy_start(e).isoformat()
        prev = cal.prev_quarter_end(e).isoformat()
        _, q = cal.fiscal(e)
        if (qs, e) in idx:
            out[e] = idx[(qs, e)]
        elif q > 1 and (fys, e) in idx and (fys, prev) in idx:
            out[e] = idx[(fys, e)] - idx[(fys, prev)]
    return pd.Series(out, dtype=float).sort_index()


# ---- comps --------------------------------------------------------------------------------------
def company_comps(c: companies.Company, cf: dict, as_of: str | None = None) -> Result:
    cal = FiscalCalendar(c.fiscal_year_end_month)
    lines = c.benchmark_lines()
    tags = c.check_tags()
    concepts = {x for v in lines.values() for x in v} | set(tags.values())
    kn = known(load_facts(cf, concepts), as_of)
    present = set(kn["concept"])
    notes = list(c.notes())

    q = {}
    for line, cs in lines.items():
        parts = {}
        for concept in cs:
            if concept not in present:
                if line in companies.BENCHMARK_REQUIRED and len(cs) == 1:
                    raise ValueError(f"{c.ticker}: required line {line} ({concept}) is not reported"
                                     f"{' by ' + as_of if as_of else ''}")
                notes.append(f"{concept} is not reported; treated as 0 in {line}")
                continue
            parts[concept] = quarterly(kn, concept, cal)
        if not parts:
            if line in companies.BENCHMARK_REQUIRED:
                raise ValueError(f"{c.ticker}: no concept for required line {line} is reported")
            q[line] = None
            continue
        if len(parts) == 1:
            q[line] = next(iter(parts.values()))
            continue
        # a summed line (e.g. capex + capitalized software): a component the company stops
        # reporting counts as 0 for those quarters, and the note says so with its last value
        frame = pd.DataFrame(parts).sort_index()
        for concept in frame.columns:
            gap = frame.index[frame[concept].isna() & (frame.index > frame[concept].first_valid_index())]
            if len(gap):
                last = frame[concept].dropna().iloc[-1]
                notes.append(f"{concept.split(':')[1]} not reported for {len(gap)} quarter(s) from {gap[0]} "
                             f"(last reported value ${last:,.0f}); counted as 0 in {line}")
        q[line] = frame.fillna(0.0).sum(axis=1)
    if "amortization" not in lines and not any("amortization" in n.lower() for n in notes):
        notes.append("no amortization concept configured; treated as 0")
    qdf = pd.DataFrame({k: v for k, v in q.items() if v is not None}).sort_index()

    # latest quarter with every required line, and the eight quarters needed for TTM and its growth
    required = [k for k, v in q.items() if v is not None]
    complete = qdf.dropna(subset=required)
    ends = list(complete.index)
    if len(ends) < 1:
        raise ValueError(f"{c.ticker}: no quarter has every required line as of {as_of or 'latest'}")
    last = ends[-1]
    window = cal.quarter_ends(cal.prev_quarter_end(cal.prev_quarter_end(cal.prev_quarter_end(
        date.fromisoformat(last)))), last)
    prior = [cal.prev_quarter_end(cal.prev_quarter_end(cal.prev_quarter_end(cal.prev_quarter_end(d)))) for d in window]
    win, pri = [d.isoformat() for d in window], [d.isoformat() for d in prior]
    missing = [d for d in win if d not in complete.index]
    if missing:
        raise ValueError(f"{c.ticker}: trailing quarters incomplete: {missing}")

    ttm = {}
    for line in TTM_LINES:
        ttm[line] = float(qdf.loc[win, line].sum()) if line in qdf else 0.0
    ttm["revenue_prior"] = float(qdf.loc[pri, "revenue"].sum()) if all(d in qdf.index for d in pri) else float("nan")

    r = ttm["revenue"]
    row = {
        "ticker": c.ticker, "name": c.name, "fiscal_label": cal.label(last), "period_end": last,
        "revenue_ttm": r, "revenue_growth_yoy": r / ttm["revenue_prior"] - 1,
        "gross_margin": 1 - ttm["cost_of_revenue"] / r,
        "rnd_pct": ttm["rnd"] / r, "snm_pct": ttm["snm"] / r, "gna_pct": ttm["gna"] / r, "sbc_pct": ttm["sbc"] / r,
        "gaap_op_margin": ttm["operating_income"] / r,
        "non_gaap_op_margin_approx": (ttm["operating_income"] + ttm["sbc"] + ttm["amortization"]) / r,
        "fcf_margin": (ttm["cfo"] - ttm["capex"]) / r,
    }
    row["rule_of_40"] = row["revenue_growth_yoy"] + row["fcf_margin"]
    res = Result(row, ttm, qdf, notes=notes)
    res.checks = reconcile(c, kn, qdf, cal, win, tags)
    # the row must describe the latest quarter the company had filed by the as-of date
    filed = kn[kn["concept"] == lines["revenue"][0]]
    latest_filed = max(e for e in filed["report_end"] if cal.is_quarter_end(e))
    res.checks.append((f"{c.ticker}: comp row is the latest filed quarter", last == latest_filed,
                       f"row {cal.label(last)} ({last}); latest filed {cal.label(latest_filed)} ({latest_filed})"))
    return res


def reconcile(c: companies.Company, kn: pd.DataFrame, qdf: pd.DataFrame, cal: FiscalCalendar,
              win: list[str], tags: dict[str, str]) -> list[tuple[str, bool, str]]:
    out = []
    # 1. derived quarters of the latest complete fiscal year sum to the 10-K figure
    lines = c.benchmark_lines()
    fy_ends = [e for e in qdf.index if cal.fiscal(e)[1] == 4]
    for fy_end in reversed(fy_ends):
        fys = cal.fy_start(fy_end).isoformat()
        q4 = date.fromisoformat(fy_end)
        qs = [d.isoformat() for d in (cal.prev_quarter_end(cal.prev_quarter_end(cal.prev_quarter_end(q4))),
                                      cal.prev_quarter_end(cal.prev_quarter_end(q4)), cal.prev_quarter_end(q4), q4)]
        if not all(d in qdf.index for d in qs):
            continue
        worst = 0.0
        for line in ("revenue", "operating_income", "cfo"):
            annual = 0.0
            for concept in lines[line]:
                hit = kn[(kn["concept"] == concept) & (kn["start"] == fys) & (kn["end"] == fy_end)]
                annual += float(hit["value"].iloc[0]) if len(hit) else float("nan")
            worst = max(worst, abs(float(qdf.loc[qs, line].sum()) - annual))
        out.append((f"{c.ticker}: quarters sum to the 10-K year ({cal.label(fy_end)[3:]})", worst <= 0.5,
                    f"max |diff| ${worst:,.0f} (revenue, operating income, CFO)"))
        break
    # 2. revenue - cost of revenue = GrossProfit where the company tags it
    if "gross_profit" in tags:
        gp = quarterly(kn, tags["gross_profit"], cal)
        common = [d for d in win if d in gp.index]
        if common:
            d = (qdf.loc[common, "revenue"] - qdf.loc[common, "cost_of_revenue"] - gp.loc[common]).abs().max()
            out.append((f"{c.ticker}: revenue - cost of revenue = GrossProfit tag", d <= 0.5, f"max |diff| ${d:,.0f}"))
    # 3. R&D + S&M + G&A versus total operating expenses (informational: other lines may exist)
    if "opex_total" in tags:
        ox = quarterly(kn, tags["opex_total"], cal)
        common = [d for d in win if d in ox.index]
        if common:
            d = (qdf.loc[common, ["rnd", "snm", "gna"]].sum(axis=1) - ox.loc[common]).abs().max()
            out.append((f"{c.ticker}: R&D + S&M + G&A = OperatingExpenses tag", d <= 0.5, f"max |diff| ${d:,.0f}"))
    return out


def comps(tickers: list[str], as_of: str | None = None, refresh: bool = False, raw_dir: Path = RAW,
          offline: dict[str, dict] | None = None) -> tuple[pd.DataFrame, dict[str, Result]]:
    """Comp table for `tickers`. `offline` maps ticker -> companyfacts dict (tests, no network)."""
    results = {}
    for t in tickers:
        c = companies.load(t)
        if "benchmark" not in c.raw:
            raise companies.ConfigError(f"{t}: config has no `benchmark` section")
        cf = offline[c.ticker] if offline else json.loads(fetch(c, refresh, raw_dir).read_text())
        results[c.ticker] = company_comps(c, cf, as_of)
    table = pd.DataFrame([r.row for r in results.values()])[list(COLUMNS)]
    return table, results


# ---- output -------------------------------------------------------------------------------------
def format_table(table: pd.DataFrame) -> str:
    t = table.copy()
    t["revenue_ttm"] = t["revenue_ttm"].map(lambda v: f"${v / 1e6:,.0f}M")
    for col in COLUMNS:
        if col not in ("ticker", "name", "fiscal_label", "period_end", "revenue_ttm"):
            t[col] = t[col].map(lambda v: f"{v * 100:.1f}%")
    return t.drop(columns=["name"]).to_string(index=False)


def write(table: pd.DataFrame, results: dict[str, Result], out: Path) -> Path:
    if out.suffix == ".csv":
        table.to_csv(out, index=False)
    elif out.suffix == ".parquet":
        import duckdb
        duckdb.connect().register("t", table).execute(f"COPY t TO '{out}' (FORMAT parquet)")
    elif out.suffix == ".xlsx":
        write_xlsx(table, results, out)
    else:
        raise ValueError("--out must end in .csv, .parquet or .xlsx")
    return out


INPUT_COLS = ["revenue", "revenue_prior", "cost_of_revenue", "rnd", "snm", "gna", "operating_income", "sbc",
              "amortization", "cfo", "capex"]
FORMULAS = {  # ratio column -> formula over the input columns of the same row ({col} = that cell)
    "revenue_growth_yoy": "={revenue}/{revenue_prior}-1",
    "gross_margin": "=1-{cost_of_revenue}/{revenue}",
    "rnd_pct": "={rnd}/{revenue}", "snm_pct": "={snm}/{revenue}", "gna_pct": "={gna}/{revenue}",
    "sbc_pct": "={sbc}/{revenue}", "gaap_op_margin": "={operating_income}/{revenue}",
    "non_gaap_op_margin_approx": "=({operating_income}+{sbc}+{amortization})/{revenue}",
    "fcf_margin": "=({cfo}-{capex})/{revenue}",
    "rule_of_40": "={revenue_growth_yoy}+{fcf_margin}",
}


def write_xlsx(table: pd.DataFrame, results: dict[str, Result], out: Path) -> Path:
    """Peer_Comps sheet: TTM inputs as values (blue), every ratio as a live formula."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Peer_Comps"
    ws["A1"] = "Public peer comps: trailing four fiscal quarters from SEC XBRL (each company's own fiscal calendar)"
    ws["A1"].font = Font(bold=True, size=13)
    head = ["ticker", "name", "fiscal_label", "period_end", *[f"{c}_ttm" for c in INPUT_COLS], *FORMULAS]
    col = {h: get_column_letter(i + 1) for i, h in enumerate(head)}
    for i, h in enumerate(head):
        c = ws.cell(3, i + 1, COLUMNS.get(h, h.replace("_ttm", " (TTM, $)")))
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    for r, res in enumerate(results.values(), start=4):
        for h in ("ticker", "name", "fiscal_label", "period_end"):
            ws[f"{col[h]}{r}"] = res.row[h]
        for c in INPUT_COLS:
            cell = ws[f"{col[c + '_ttm']}{r}"]
            cell.value, cell.font, cell.number_format = res.ttm[c], Font(color="0000FF"), '$#,##0'
        for name, tmpl in FORMULAS.items():
            f = tmpl
            for c in INPUT_COLS:
                f = f.replace("{" + c + "}", f"{col[c + '_ttm']}{r}")
            for c in FORMULAS:
                f = f.replace("{" + c + "}", f"{col[c]}{r}")
            cell = ws[f"{col[name]}{r}"]
            cell.value, cell.number_format = f, "0.0%"
    r = len(results) + 5
    ws.cell(r, 1, "Notes").font = Font(bold=True)
    for t, res in results.items():
        for n in res.notes:
            r += 1
            ws.cell(r, 1, f"{t}: {n}")
    for i in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.column_dimensions["B"].width = 24
    wb.save(out)
    return out
